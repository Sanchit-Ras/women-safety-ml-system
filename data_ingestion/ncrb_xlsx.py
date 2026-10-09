"""Reader for NCRB's district-wise workbooks (Crime in India -> Additional Tables -> District Wise Reports).

The 2015 workbook has one data sheet. Its shape:

    row 1   title
    row 2   header: SL | District | YEAR | Rape | Custodial Rape | ... (60 crime heads, merged-cell gaps empty)
    row 3   'State: Andhra Pradesh'          <- a marker: the rows below belong to this state
    row 4+  1 | Anantapur | 2015 | 32 | ...  <- one police district per row
            ...
            n | Total District(s) | ...      <- the state's sum: skipped, never a unit
    ...
    last    'Source: Crime in India'

Totals sit beside their parts ('Kidnapping & Abduction_Total' followed by five sub-heads): the reader keeps
every column as it is; deciding which columns count is the crosswalk's job, not the reader's.
Every emitted row carries ``source_row`` = the Excel row number, so a record can be found in the sheet.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import pandas as pd

HEADER_FIRST_CELL = "SL"
STATE_MARKER = re.compile(r"^\s*(?:State|UT)\s*:\s*(.+?)\s*$", re.I)
KEY_LABELS = ("SL", "District", "YEAR")                # header labels that are not crime heads
ROW_KEYS = ("state_name", "district", "source_row")  # keys the reader adds to every row


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_district_rows(path: Path | str, sheet: str | None = None) -> list[dict]:
    """One dict per district row: state_name, district, source_row, YEAR and every crime head -> cell text."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    header: list[str] | None = None
    state: str | None = None
    out: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        first = _text(row[0]) if row else ""
        if header is None:
            if first == HEADER_FIRST_CELL:
                header = [_text(c) for c in row]
            continue
        marker = STATE_MARKER.match(first)
        if marker:
            state = marker.group(1)
            continue
        name = _text(row[1]) if len(row) > 1 else ""
        if not name or name.lower().startswith("total") or first.lower().startswith("source"):
            continue
        if state is None:
            raise ValueError(f"{path}: row {i} ('{name}') comes before any 'State:' marker")
        rec = {"state_name": state, "district": name, "source_row": i}
        for label, value in zip(header, row):
            if label and label not in ("SL", "District"):
                rec[label] = _text(value)
        out.append(rec)
    wb.close()
    if header is None:
        raise ValueError(f"{path}: no header row starting with '{HEADER_FIRST_CELL}'")
    return out


def crime_columns(path: Path | str, sheet: str | None = None) -> list[str]:
    """The crime heads in sheet order: the header minus the key labels."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    for row in ws.iter_rows(values_only=True):
        if row and _text(row[0]) == HEADER_FIRST_CELL:
            wb.close()
            return [_text(c) for c in row if _text(c) and _text(c) not in KEY_LABELS]
    wb.close()
    raise ValueError(f"{path}: no header row starting with '{HEADER_FIRST_CELL}'")


def frame(path: Path | str, sheet: str | None = None) -> pd.DataFrame:
    """The district rows as a table with numeric crime columns and an integer ``year`` column."""
    rows = read_district_rows(path, sheet)
    df = pd.DataFrame(rows)
    df["year"] = pd.to_numeric(df["YEAR"], errors="coerce").astype("Int64")
    for col in df.columns:
        if col not in ROW_KEYS and col not in KEY_LABELS and col != "year":
            df[col] = pd.to_numeric(df[col], errors="coerce")
    return df
